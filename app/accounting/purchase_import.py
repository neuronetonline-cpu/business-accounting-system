from datetime import date, datetime
from pathlib import Path
import re

from openpyxl import load_workbook, Workbook

from app.database import get_connection
from app.accounting.inventory import create_purchase
from app.accounting.audit import audit


def _norm(v):
    return re.sub(r'[^a-z0-9]+', ' ', str(v or '').strip().lower()).strip()


def _num(v, default=0.0):
    if v is None or str(v).strip() == '':
        return float(default)
    s = str(v).replace(',', '').replace('Rs.', '').replace('LKR', '').strip()
    return float(s or default)


def _date(v):
    if not v:
        return date.today().isoformat()
    if isinstance(v, datetime):
        return v.date().isoformat()
    if hasattr(v, 'isoformat') and not isinstance(v, str):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y', '%b %d, %Y'):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f'Invalid date: {s}')


def _find_index(headers, aliases):
    for alias in aliases:
        if alias in headers:
            return headers[alias]
    return None


def read_purchase_workbook(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        raise ValueError('Empty Excel workbook.')
    headers = {_norm(v): i for i, v in enumerate(header) if v not in (None, '')}

    idx = {
        'date': _find_index(headers, ['date', 'purchase date']),
        'bill': _find_index(headers, ['bill no', 'bill number', 'invoice no', 'invoice number', 'reference']),
        'sku': _find_index(headers, ['sku', 'barcode', 'product code', 'item code']),
        'name': _find_index(headers, ['product name', 'product', 'name', 'item name', 'item']),
        'qty': _find_index(headers, ['qty', 'quantity']),
        'cost': _find_index(headers, ['unit cost', 'cost', 'cost price', 'purchase price', 'unit price']),
        'selling': _find_index(headers, ['selling price', 'sale price', 'selling']),
        'supplier': _find_index(headers, ['supplier', 'supplier name', 'vendor']),
        'payment': _find_index(headers, ['payment type', 'payment', 'pay type']),
        'paid': _find_index(headers, ['paid', 'amount paid', 'paid amount']),
        'category': _find_index(headers, ['category']),
        'brand': _find_index(headers, ['brand']),
        'unit': _find_index(headers, ['unit']),
        'reorder': _find_index(headers, ['reorder level', 'reorder']),
    }
    if idx['bill'] is None or idx['qty'] is None or idx['cost'] is None or (idx['sku'] is None and idx['name'] is None):
        raise ValueError('Purchase Excel must contain Bill No, Product/SKU, Qty and Unit Cost columns.')

    def val(row, key, default=''):
        i = idx[key]
        return row[i] if i is not None and i < len(row) and row[i] is not None else default

    records = []
    for n, row in enumerate(rows, start=2):
        if not row or all(v in (None, '') for v in row):
            continue
        try:
            bill = str(val(row, 'bill')).strip()
            sku = str(val(row, 'sku')).strip()
            name = str(val(row, 'name')).strip()
            if not bill:
                raise ValueError('Bill No is required')
            if not sku and not name:
                raise ValueError('SKU or Product Name is required')
            qty = _num(val(row, 'qty'))
            cost = _num(val(row, 'cost'))
            if qty <= 0:
                raise ValueError('Qty must be greater than zero')
            if cost < 0:
                raise ValueError('Unit Cost cannot be negative')
            records.append({
                'row': n, 'date': _date(val(row, 'date')), 'bill': bill, 'sku': sku,
                'name': name, 'qty': qty, 'cost': cost, 'selling': _num(val(row, 'selling')),
                'supplier': str(val(row, 'supplier')).strip(),
                'payment': str(val(row, 'payment', 'Cash')).strip() or 'Cash',
                'paid': _num(val(row, 'paid')),
                'category': str(val(row, 'category')).strip(), 'brand': str(val(row, 'brand')).strip(),
                'unit': str(val(row, 'unit', 'pcs')).strip() or 'pcs', 'reorder': _num(val(row, 'reorder')),
            })
        except Exception as exc:
            raise ValueError(f'Row {n}: {exc}')
    if not records:
        raise ValueError('No purchase rows found.')
    return records


def _resolve_supplier(con, name):
    if not name or name.lower() in ('none', 'walk-in', 'walk in'):
        return None
    row = con.execute('SELECT id FROM suppliers WHERE lower(name)=lower(?) AND active=1', (name,)).fetchone()
    if row:
        return row['id']
    raise ValueError(f'Supplier not found: {name}. Add the supplier first.')


def _resolve_or_create_product(con, rec):
    product = None
    if rec['sku']:
        product = con.execute('SELECT * FROM products WHERE sku=?', (rec['sku'],)).fetchone()
    if not product and rec['name']:
        product = con.execute('SELECT * FROM products WHERE lower(name)=lower(?)', (rec['name'],)).fetchone()
    if product:
        return product['id']
    if not rec['name']:
        raise ValueError(f"Row {rec['row']}: Product not found and Product Name is empty.")
    con.execute('''INSERT INTO products(sku,name,category,brand,unit,cost_price,selling_price,reorder_level)
                   VALUES(?,?,?,?,?,?,?,?)''',
                (rec['sku'] or None, rec['name'], rec['category'], rec['brand'], rec['unit'],
                 rec['cost'], rec['selling'], rec['reorder']))
    return con.execute('SELECT last_insert_rowid()').fetchone()[0]


def import_purchase_excel(path):
    records = read_purchase_workbook(path)
    con = get_connection()
    try:
        existing = []
        for bill in sorted({r['bill'] for r in records}):
            if con.execute('SELECT COUNT(*) n FROM purchases WHERE bill_no=?', (bill,)).fetchone()['n']:
                existing.append(bill)
        if existing:
            raise ValueError('These Bill/Invoice numbers already exist: ' + ', '.join(existing[:20]))

        grouped = {}
        for rec in records:
            grouped.setdefault(rec['bill'], []).append(rec)
        preview = []
        for bill, rows in grouped.items():
            dates = {r['date'] for r in rows}
            suppliers = {r['supplier'] for r in rows if r['supplier']}
            payments = {r['payment'] for r in rows if r['payment']}
            if len(dates) != 1:
                raise ValueError(f'Bill {bill}: all rows must have the same Date.')
            if len(suppliers) > 1:
                raise ValueError(f'Bill {bill}: all rows must have the same Supplier.')
            if len(payments) > 1:
                raise ValueError(f'Bill {bill}: all rows must have the same Payment Type.')
            paid_values = [r['paid'] for r in rows if r['paid'] > 0]
            if len(paid_values) > 1 and any(abs(x-paid_values[0]) > 0.005 for x in paid_values[1:]):
                raise ValueError(f'Bill {bill}: Paid amount should be entered once for the bill, not differently on each line.')
            subtotal = round(sum(r['qty'] * r['cost'] for r in rows), 2)
            paid = paid_values[0] if paid_values else 0
            if paid > subtotal:
                raise ValueError(f'Bill {bill}: Paid amount exceeds purchase total.')
            preview.append((bill, rows[0]['date'], rows[0]['supplier'] or 'None', rows[0]['payment'] or 'Cash', len(rows), subtotal, paid))
        return records, grouped, preview
    finally:
        con.close()


def commit_purchase_import(records, grouped):
    imported = 0
    created_products = 0
    for bill, rows in grouped.items():
        con = get_connection()
        try:
            before = con.execute('SELECT COUNT(*) n FROM products').fetchone()['n']
            supplier_id = _resolve_supplier(con, rows[0]['supplier'])
            items = []
            for rec in rows:
                pid = _resolve_or_create_product(con, rec)
                items.append((pid, rec['qty'], rec['cost']))
            con.commit()
            after = con.execute('SELECT COUNT(*) n FROM products').fetchone()['n']
            created_products += max(0, after - before)
        finally:
            con.close()
        paid = next((r['paid'] for r in rows if r['paid'] > 0), 0)
        create_purchase(rows[0]['date'], bill, supplier_id, rows[0]['payment'] or 'Cash', items, paid)
        imported += 1
    audit('PURCHASE_EXCEL_IMPORT', str(Path(records[0]['row']) if records else ''),
          f'Imported {imported} purchase bills from Excel; {created_products} products created')
    return imported, created_products


def create_purchase_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Purchases'
    headers = ['Date','Bill No','SKU / Barcode','Product Name','Qty','Unit Cost','Selling Price','Supplier','Payment Type','Paid','Category','Brand','Unit','Reorder Level']
    ws.append(headers)
    ws.append(['2026-09-28','PUR-001','SSD-001','Example SSD',2,20000,25000,'TEST SUPPLIER','Cash',40000,'SSD','ExampleBrand','pcs',5])
    ws.append(['2026-09-28','PUR-002','RAM-001','Example RAM',5,8000,10000,'TEST SUPPLIER','Credit',0,'RAM','ExampleBrand','pcs',2])
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    ws.freeze_panes = 'A2'
    wb.save(path)
