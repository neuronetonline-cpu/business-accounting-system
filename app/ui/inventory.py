import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from app.database import get_connection
from app.accounting.inventory import add_product, product_balance
from app.accounting.audit import audit
from app.ui.theme import COLORS, FONT
from app.ui.widgets import tree_with_scrollbars, Card

class InventoryFrame(tk.Frame):
    def __init__(self, master):
        super().__init__(master,bg=COLORS['bg']); self.fields={}; self.build()
    def build(self):
        tk.Label(self,text='Products & Inventory',bg=COLORS['bg'],fg=COLORS['text'],font=(FONT,22,'bold')).pack(anchor='w')
        tk.Label(self,text='Add products individually or import many products from an Excel workbook.',bg=COLORS['bg'],fg=COLORS['muted'],font=(FONT,9)).pack(anchor='w',pady=(2,12))
        top=Card(self,padx=18,pady=14); top.pack(fill='x',pady=(0,12))
        labels=['SKU','Product Name','Category','Brand','Unit','Cost Price','Selling Price','Reorder Level']
        for i,l in enumerate(labels):
            tk.Label(top,text=l,bg='white',fg=COLORS['text'],font=(FONT,8,'bold')).grid(row=0,column=i,padx=3,sticky='w')
            e=ttk.Entry(top,width=15); e.grid(row=1,column=i,padx=3,pady=5); self.fields[l]=e
        ttk.Button(top,text='＋ ADD PRODUCT',style='Accent.TButton',command=self.add).grid(row=1,column=8,padx=7)
        ttk.Button(top,text='⇩ IMPORT EXCEL',command=self.import_excel).grid(row=1,column=9,padx=7)
        ttk.Button(top,text='⇧ TEMPLATE',command=self.template).grid(row=1,column=10,padx=7)
        box=Card(self,padx=10,pady=10); box.pack(fill='both',expand=True)
        cols=('ID','SKU','Product','Category','Brand','Unit','Cost','Selling','Qty','Stock Value','Reorder','Status')
        frame,self.tree=tree_with_scrollbars(box,cols,{'ID':55,'SKU':120,'Product':190,'Category':110,'Brand':110,'Unit':70,'Cost':110,'Selling':110,'Qty':90,'Stock Value':130,'Reorder':90,'Status':90},height=18)
        frame.pack(fill='both',expand=True)
        actions=tk.Frame(self,bg=COLORS['bg']); actions.pack(fill='x',pady=10)
        ttk.Button(actions,text='EDIT SELECTED',command=self.edit_selected).pack(side='right')
        ttk.Button(actions,text='DEACTIVATE SELECTED',command=self.deactivate).pack(side='right',padx=8)
        ttk.Button(actions,text='DELETE SELECTED',command=self.delete_selected).pack(side='right',padx=8)
        ttk.Button(actions,text='REFRESH',command=self.refresh).pack(side='right')
        self.refresh()
    def selected(self):
        sel=self.tree.selection()
        if not sel:return None
        pid=int(self.tree.item(sel[0])['values'][0]); con=get_connection(); r=con.execute('SELECT * FROM products WHERE id=?',(pid,)).fetchone(); con.close(); return r
    def edit_selected(self):
        r=self.selected()
        if not r:return
        win=tk.Toplevel(self); win.title('Edit Product'); win.geometry('620x520'); win.configure(bg=COLORS['bg'])
        box=Card(win,padx=18,pady=18); box.pack(fill='both',expand=True,padx=15,pady=15); fields={}
        labels=['SKU','Product Name','Category','Brand','Unit','Cost Price','Selling Price','Reorder Level']; keys=['sku','name','category','brand','unit','cost_price','selling_price','reorder_level']
        for i,(lab,key) in enumerate(zip(labels,keys)):
            tk.Label(box,text=lab,bg='white',font=(FONT,9,'bold')).grid(row=i,column=0,sticky='w',pady=6);e=ttk.Entry(box,width=45);e.insert(0,str(r[key] or ''));e.grid(row=i,column=1,padx=12,pady=6);fields[key]=e
        def save():
            try:
                vals=[fields[k].get().strip() for k in keys]; vals[5]=float(vals[5].replace(',','') or 0); vals[6]=float(vals[6].replace(',','') or 0); vals[7]=float(vals[7].replace(',','') or 0)
                if not vals[1]:raise ValueError('Product Name is required.')
                con=get_connection();dup=con.execute('SELECT id FROM products WHERE sku=? AND id<>?',(vals[0] or None,r['id'])).fetchone() if vals[0] else None
                if dup:raise ValueError('SKU already exists.')
                con.execute('UPDATE products SET sku=?,name=?,category=?,brand=?,unit=?,cost_price=?,selling_price=?,reorder_level=? WHERE id=?',(*vals,r['id']));con.commit();con.close();audit('PRODUCT_EDIT',str(r['id']),f'Product edited: {vals[1]}');messagebox.showinfo('Saved','Product updated successfully.');win.destroy();self.refresh()
            except Exception as e:messagebox.showerror('Edit Error',str(e))
        ttk.Button(win,text='SAVE CHANGES',style='Accent.TButton',command=save).pack(side='right',padx=20,pady=(0,15));ttk.Button(win,text='CANCEL',command=win.destroy).pack(side='right',padx=8,pady=(0,15))
    def delete_selected(self):
        r=self.selected()
        if not r:return
        try:
            con=get_connection();linked=0
            for table in ('stock_movements','sale_items','purchase_items'):
                linked += con.execute(f'SELECT COUNT(*) n FROM {table} WHERE product_id=?',(r['id'],)).fetchone()['n']
            con.close()
            if linked:
                messagebox.showinfo('Cannot Delete','This product has stock or transaction history. Use DEACTIVATE instead so history is preserved.');return
            if not messagebox.askyesno('Delete',f'Delete {r["name"]}? This cannot be undone.'):return
            con=get_connection();con.execute('DELETE FROM products WHERE id=?',(r['id'],));con.commit();con.close();audit('PRODUCT_DELETE',str(r['id']),f'Product deleted: {r["name"]}');self.refresh()
        except Exception as e:messagebox.showerror('Delete Error',str(e))
    def deactivate(self):
        r=self.selected()
        if not r:return
        if not messagebox.askyesno('Deactivate',f'Deactivate {r["name"]}?\n\nStock and transaction history will be preserved.'):return
        con=get_connection();con.execute('UPDATE products SET active=0 WHERE id=?',(r['id'],));con.commit();con.close();audit('PRODUCT_DEACTIVATE',str(r['id']),f'Product deactivated: {r["name"]}');self.refresh()
    def add(self):
        try:
            add_product(self.fields['SKU'].get().strip(),self.fields['Product Name'].get().strip(),self.fields['Category'].get().strip(),self.fields['Brand'].get().strip(),self.fields['Unit'].get().strip() or 'pcs',float(self.fields['Cost Price'].get().replace(',','') or 0),float(self.fields['Selling Price'].get().replace(',','') or 0),float(self.fields['Reorder Level'].get().replace(',','') or 0))
            audit('PRODUCT_ADD',self.fields['SKU'].get().strip(),f"Product added: {self.fields['Product Name'].get().strip()}")
            for e in self.fields.values():e.delete(0,'end')
            self.refresh(); messagebox.showinfo('Saved','Product added successfully.')
        except Exception as e: messagebox.showerror('Error',str(e))
    def import_excel(self):
        path=filedialog.askopenfilename(title='Select Product / POS Stock Excel File',filetypes=[('Excel Workbook','*.xlsx'),('Excel 97-2003','*.xls'),('All Files','*.*')])
        if not path:return
        try:
            from openpyxl import load_workbook
            from datetime import datetime, date
            import re
            wb=load_workbook(path,data_only=True,read_only=True); ws=wb.active
            rows_iter=ws.iter_rows(values_only=True); header=next(rows_iter,None)
            if not header: raise ValueError('Empty Excel workbook.')
            raw_headers=[str(c or '').strip().lower() for c in header]
            norm_headers={re.sub(r'[^a-z0-9]+',' ',h).strip():i for i,h in enumerate(raw_headers)}
            pos_required={'name','barcode','cost price','sale price','total stock'}
            if pos_required.issubset(set(norm_headers.keys())):
                def idx(*names):
                    for n in names:
                        if n in norm_headers:return norm_headers[n]
                    return None
                name_i,sku_i,cost_i,sale_i,qty_i,date_i=idx('name'),idx('barcode'),idx('cost price'),idx('sale price'),idx('total stock'),idx('date')
                data=[]; dates=[]
                for n,row in enumerate(rows_iter,start=2):
                    if not row or all(v in (None,'') for v in row): continue
                    name=str(row[name_i] or '').strip()
                    if not name: continue
                    sku=str(row[sku_i] or '').strip() if sku_i is not None else ''
                    def num(v): return float(str(v or 0).replace(',','').replace('Rs.','').strip() or 0)
                    cost,sale,qty=num(row[cost_i]),num(row[sale_i]),num(row[qty_i])
                    if qty < 0: raise ValueError(f'Negative stock at Excel row {n}: {name}')
                    if date_i is not None and row[date_i]:
                        try:
                            dv=row[date_i]
                            d=dv.date() if isinstance(dv,(datetime,date)) else datetime.strptime(str(dv).strip(),'%b %d, %Y').date()
                            dates.append(d)
                        except Exception: pass
                    data.append((sku,name,cost,sale,qty))
                if not data: raise ValueError('No POS stock rows found.')
                opening_date=max(dates).isoformat() if dates else date.today().isoformat()
                if not messagebox.askyesno('POS Stock Report Detected',f'This Excel file matches the POS stock report format.\n\nProducts found: {len(data)}\nOpening stock date: {opening_date}\n\nImport products and use TOTAL STOCK as opening stock for that date?'): return
                con=get_connection(); created=updated=stock_rows=0
                try:
                    for sku,name,cost,sale,qty in data:
                        product=None
                        if sku: product=con.execute('SELECT * FROM products WHERE sku=?',(sku,)).fetchone()
                        if not product: product=con.execute('SELECT * FROM products WHERE lower(name)=lower(?)',(name,)).fetchone()
                        if product:
                            pid=product['id']; con.execute('UPDATE products SET sku=CASE WHEN ?<>'' THEN ? ELSE sku END,name=?,cost_price=?,selling_price=? WHERE id=?',(sku,sku,name,cost,sale,pid)); updated+=1
                        else:
                            con.execute('INSERT INTO products(sku,name,cost_price,selling_price) VALUES(?,?,?,?)',(sku or None,name,cost,sale)); pid=con.execute('SELECT last_insert_rowid()').fetchone()[0]; created+=1
                        con.execute("DELETE FROM stock_movements WHERE product_id=? AND movement_date=? AND movement_type='OPENING'",(pid,opening_date))
                        if qty:
                            con.execute('INSERT INTO stock_movements(product_id,movement_date,reference,movement_type,qty,unit_cost,total_cost) VALUES(?,?,?,?,?,?,?)',(pid,opening_date,'POS-STOCK-IMPORT','OPENING',qty,cost,round(qty*cost,2))); stock_rows+=1
                    con.commit()
                except Exception: con.rollback(); raise
                finally: con.close()
                audit('PRODUCT_POS_STOCK_IMPORT',path,f'POS stock import: {len(data)} products, {stock_rows} opening stock rows, date {opening_date}')
                self.refresh(); messagebox.showinfo('POS Stock Import',f'POS stock report imported successfully.\n\nProducts: {len(data)}\nNew products: {created}\nUpdated products: {updated}\nOpening stock rows: {stock_rows}\nOpening date: {opening_date}')
                return
            aliases={'name':'product name','product':'product name','cost':'cost price','selling':'selling price','reorder':'reorder level'}
            headers=[aliases.get(h,h) for h in raw_headers]; idxmap={h:i for i,h in enumerate(headers)}
            if 'sku' not in idxmap or 'product name' not in idxmap: raise ValueError('Excel must contain SKU and Product Name columns, or use the POS Stock Report format (NAME, BARCODE, COST PRICE, SALE PRICE, TOTAL STOCK).')
            rows=[]; errors=[]
            for n,row in enumerate(rows_iter,start=2):
                if not any(v not in (None,'') for v in row): continue
                def val(key,default=''):
                    i=idxmap.get(key); return row[i] if i is not None and i<len(row) and row[i] is not None else default
                try: rows.append((str(val('sku')).strip(),str(val('product name')).strip(),str(val('category')).strip(),str(val('brand')).strip(),str(val('unit','pcs')).strip() or 'pcs',float(val('cost price',0) or 0),float(val('selling price',0) or 0),float(val('reorder level',0) or 0)))
                except Exception as ex: errors.append(f'Row {n}: {ex}')
            if not rows: raise ValueError('No product rows found.')
            con=get_connection()
            try: con.executemany('INSERT INTO products(sku,name,category,brand,unit,cost_price,selling_price,reorder_level) VALUES(?,?,?,?,?,?,?,?)',rows); con.commit()
            except Exception: con.rollback(); raise
            finally: con.close()
            audit('PRODUCT_IMPORT',path,f'Imported {len(rows)} product(s) from Excel'); self.refresh(); msg=f'{len(rows)} product(s) imported successfully.'
            if errors: msg+='\n\nSkipped rows:\n'+'\n'.join(errors[:10])
            messagebox.showinfo('Excel Import',msg)
        except ImportError: messagebox.showerror('Excel Import','openpyxl is required. Run: py -m pip install openpyxl')
        except Exception as e: messagebox.showerror('Excel Import',str(e))
    def template(self):
        try:
            from openpyxl import Workbook
            path=filedialog.asksaveasfilename(title='Save Product Excel Template',defaultextension='.xlsx',filetypes=[('Excel Workbook','*.xlsx')],initialfile='products_import_template.xlsx')
            if not path:return
            wb=Workbook();ws=wb.active;ws.title='Products'
            headers=['SKU','Product Name','Category','Brand','Unit','Cost Price','Selling Price','Reorder Level']; ws.append(headers)
            ws.append(['SSD-001','Example SSD','SSD','ExampleBrand','pcs',20000,25000,5])
            wb.save(path); messagebox.showinfo('Template','Excel template created.')
        except Exception as e:messagebox.showerror('Template Error',str(e))
    def refresh(self):
        for x in self.tree.get_children(): self.tree.delete(x)
        con=get_connection(); rows=con.execute('SELECT * FROM products WHERE active=1 ORDER BY name').fetchall()
        for r in rows:
            q,c=product_balance(r['id']); status='LOW' if q<=r['reorder_level'] else 'OK'
            self.tree.insert('','end',values=(r['id'],r['sku'] or '',r['name'],r['category'] or '',r['brand'] or '',r['unit'] or 'pcs',f"{r['cost_price']:,.2f}",f"{r['selling_price']:,.2f}",f"{q:,.2f}",f"{c:,.2f}",f"{r['reorder_level']:,.2f}",status))
        con.close()
